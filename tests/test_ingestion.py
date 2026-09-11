"""
Tests for core/ingestion/pdf_extractor.py

Run with (from the project root):
    pytest
"""

import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pytest
from pypdf import PdfWriter

from core.ingestion.pdf_extractor import (
    extract_text_from_pdf,
    PDFExtractionError,
    ExtractedDocument,
    PageText,
)


@pytest.fixture
def blank_pdf(tmp_path):
    """
    Creates a real, valid, minimal 2-page PDF on disk using pypdf itself
    (no extra library needed). Pages are blank, so extracted text will
    be empty -- this test checks that the pipeline doesn't crash and
    returns the correct structure, not that specific text is present.
    """
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)

    pdf_path = tmp_path / "blank_test.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)

    return str(pdf_path)


def test_extract_missing_file_raises():
    with pytest.raises(PDFExtractionError):
        extract_text_from_pdf("this_file_does_not_exist.pdf")


def test_extract_non_pdf_raises(tmp_path):
    fake_file = tmp_path / "notes.txt"
    fake_file.write_text("hello")
    with pytest.raises(PDFExtractionError):
        extract_text_from_pdf(str(fake_file))


def test_extract_blank_pdf_structure(blank_pdf):
    document = extract_text_from_pdf(blank_pdf)
    assert isinstance(document, ExtractedDocument)
    assert document.page_count == 2
    assert document.pages[0].page_number == 1
    assert document.pages[1].page_number == 2
    # Blank pages -> empty text, but extraction must not crash
    assert document.pages[0].text == ""


def test_full_text_property_joins_pages():
    document = ExtractedDocument(
        source_path="fake.pdf",
        pages=[PageText(page_number=1, text="Page one."), PageText(page_number=2, text="Page two.")],
    )
    assert "Page one." in document.full_text
    assert "Page two." in document.full_text


# ---------- DOCX extraction ----------

def test_extract_docx_structure(tmp_path):
    from docx import Document as DocxDocument
    from core.ingestion.docx_extractor import extract_text_from_docx, DocxExtractionError

    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("This is a DocMind test document.")
    doc.add_paragraph("It has multiple paragraphs of real text.")
    doc.save(str(docx_path))

    extracted = extract_text_from_docx(str(docx_path))
    assert extracted.page_count == 1  # fits in one section (< 15 paragraphs)
    assert "DocMind test document" in extracted.full_text


def test_extract_docx_missing_file_raises():
    from core.ingestion.docx_extractor import extract_text_from_docx, DocxExtractionError

    with pytest.raises(DocxExtractionError):
        extract_text_from_docx("no_such_file.docx")


def test_extract_docx_empty_document_raises(tmp_path):
    from docx import Document as DocxDocument
    from core.ingestion.docx_extractor import extract_text_from_docx, DocxExtractionError

    docx_path = tmp_path / "empty.docx"
    DocxDocument().save(str(docx_path))  # no paragraphs added

    with pytest.raises(DocxExtractionError):
        extract_text_from_docx(str(docx_path))
