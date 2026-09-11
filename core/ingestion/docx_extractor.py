"""
DOCX text extraction module.

DOCX files don't have a native "page" concept the way PDFs do (page
breaks depend on rendering, font, and printer settings -- python-docx
cannot reliably report page numbers). To keep the rest of the pipeline
format-agnostic, we reuse the same ExtractedDocument/PageText shape as
the PDF extractor, but "page_number" here means a synthetic "section
number" (paragraphs grouped in batches). This is documented clearly so
citations from DOCX say "Section 3" instead of falsely claiming "Page 3".
"""

from pathlib import Path
from typing import List

from docx import Document as DocxDocument

from core.ingestion.pdf_extractor import ExtractedDocument, PageText, PDFExtractionError


PARAGRAPHS_PER_SECTION = 15  # groups paragraphs into pseudo-pages for citation purposes


class DocxExtractionError(Exception):
    """Raised when a DOCX file cannot be found or parsed."""


def extract_text_from_docx(file_path: str) -> ExtractedDocument:
    """
    Extract text from a DOCX file, grouped into pseudo-page "sections".

    Args:
        file_path: Path to a .docx file on disk.

    Returns:
        ExtractedDocument with synthetic page numbers (sections).

    Raises:
        DocxExtractionError: if the file doesn't exist, isn't a .docx,
            or can't be parsed.
    """
    path = Path(file_path)

    if not path.exists():
        raise DocxExtractionError(f"File not found: {file_path}")

    if path.suffix.lower() != ".docx":
        raise DocxExtractionError(f"Not a DOCX file: {file_path}")

    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocxExtractionError(f"Could not open DOCX: {exc}") from exc

    paragraphs = [p.text for p in document.paragraphs if p.text and p.text.strip()]

    if not paragraphs:
        raise DocxExtractionError(f"No extractable text found in DOCX: {file_path}")

    pages: List[PageText] = []
    for section_index in range(0, len(paragraphs), PARAGRAPHS_PER_SECTION):
        section_paragraphs = paragraphs[section_index: section_index + PARAGRAPHS_PER_SECTION]
        section_number = (section_index // PARAGRAPHS_PER_SECTION) + 1
        pages.append(
            PageText(page_number=section_number, text="\n".join(section_paragraphs))
        )

    return ExtractedDocument(source_path=str(path), pages=pages)
